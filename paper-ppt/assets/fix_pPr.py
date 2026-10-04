"""pptx 后处理 + 校验。用法：python fix_pPr.py <deck.pptx>

1) 去重 <a:p> 内多余的 <a:pPr>（pptxgenjs 多 run 段落会产生重复 pPr，
   LibreOffice 能渲染但 PowerPoint 会重排导致乱行）。
2) 用 python-pptx 重新打开校验，打印每页图片数与 notes 有无。
"""
import re
import shutil
import sys
import zipfile

src = sys.argv[1]
tmp = src + ".fixed"


def fix_p(m):
    inner = m.group(1)
    pprs = re.findall(r"<a:pPr[^>]*/>|<a:pPr[^>]*>.*?</a:pPr>", inner, flags=re.S)
    if len(pprs) <= 1:
        return m.group(0)
    first = pprs[0]
    out = inner.replace(first, "\x00M\x00", 1)
    for r in pprs[1:]:
        out = out.replace(r, "")
    return "<a:p>" + out.replace("\x00M\x00", first) + "</a:p>"


zin = zipfile.ZipFile(src)
zout = zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED)
patched = 0
for it in zin.infolist():
    data = zin.read(it.filename)
    if it.filename.startswith("ppt/slides/slide") and it.filename.endswith(".xml"):
        xml = data.decode("utf-8")
        fixed = re.sub(r"<a:p>(.*?)</a:p>", fix_p, xml, flags=re.S)
        patched += fixed != xml
        data = fixed.encode("utf-8")
    zout.writestr(it, data)
zout.close()
zin.close()
shutil.move(tmp, src)

from pptx import Presentation  # noqa: E402

prs = Presentation(src)
print(f"patched slides: {patched}, total: {len(prs.slides)}")
for i, sl in enumerate(prs.slides, 1):
    pics = sum(1 for sh in sl.shapes if sh.shape_type == 13)
    notes = "Y" if sl.has_notes_slide and sl.notes_slide.notes_text_frame.text.strip() else "N"
    print(f"s{i}: pics={pics} notes={notes}")
